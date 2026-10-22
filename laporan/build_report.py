from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).parent
REF = BASE / "reference.docx"
MD = BASE / "T9-test-execution-summary.md"
DOCX = BASE / "T9-test-execution-summary.docx"

# ---------- reference doc ----------
# Start from pandoc's default reference so table/compact styles remain available.
import subprocess
REF_BYTES = subprocess.check_output(["pandoc", "--print-default-data-file", "reference.docx"])
REF.write_bytes(REF_BYTES)
doc = Document(REF)
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.5)
sec.bottom_margin = Cm(2.5)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.5)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(6)

for name, size, color in [("Heading 1", 16, "1F3864"), ("Heading 2", 13, "2E5496"), ("Heading 3", 11.5, "2E5496")]:
    st = styles[name]
    st.font.name = "Arial"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True

# footer with page number (pandoc may preserve it; post-process also adds if missing)
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer_p.add_run("T9 — Test Execution Summary Report | CilacapTools Sprint 5 | ")
r.font.size = Pt(8)
r.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
r2 = footer_p.add_run()
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
r2._r.addnext(fld)

doc.save(REF)
print("reference saved")

# ---------- pandoc ----------
import subprocess
subprocess.run([
    "pandoc", str(MD), "-f", "markdown+pipe_tables+link_attributes", "-t", "docx",
    "-o", str(DOCX), f"--reference-doc={REF}", f"--resource-path={BASE}"
], check=True)

# ---------- post-process ----------
doc = Document(DOCX)
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.5)
sec.bottom_margin = Cm(2.5)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.5)

# header + footer (replace pandoc footer if needed)
header = sec.header.paragraphs[0]
header.text = "UTS PPL — CilacapTools Sprint 5 | T9 Test Execution Summary"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)

footer = sec.footer.paragraphs[0]
footer.text = ""
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("T9 — Test Execution Summary Report | CilacapTools Sprint 5 | halaman ")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
run2 = footer.add_run()
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
run2._r.addnext(fld)

def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), 'A6A6A6')
        borders.append(el)
    tblPr.append(borders)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

for table in doc.tables:
    table.style = "Table"
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    keep_together = len(table.rows) <= 15
    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        cant.set(qn("w:val"), "true")
        trPr.append(cant)
        if i == 0:
            set_repeat_header(row)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    if i == 0:
                        r.font.bold = True
                if keep_together and i < len(table.rows) - 1:
                    p.paragraph_format.keep_with_next = True
            if i == 0:
                shade(cell, "D9E2F3")

# Remove pandoc's alt-text paragraphs that duplicate the image titles.
ALTS = {
    "Status eksekusi Run 31 (34 test case)",
    "Hasil per priority (jumlah test case)",
    "Hasil per priority",
    "Distribusi severity 32 issue",
    "Defect density per area (bug / test case x 100%)",
}
for p in list(doc.paragraphs):
    if p.text.strip() in ALTS:
        p._element.getparent().remove(p._element)

# Column widths for the wide analytical/appendix tables.
def set_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for idx, w in enumerate(widths):
        if idx < len(table.columns):
            table.columns[idx].width = Cm(w)
        for row in table.rows:
            if idx < len(row.cells):
                row.cells[idx].width = Cm(w)

TABLE_WIDTHS = {
    "Field": [4.0, 12.0],
    "Item": [4.5, 11.5],
    "Metrik": [7.0, 4.5, 4.5],
    "Priority": [3.0, 2.0, 2.0, 2.0, 2.0, 3.0],
    "Paket fitur": [4.8, 1.2, 1.3, 1.3, 1.3, 1.4, 1.8, 1.8],
    "Area": [3.6, 1.8, 1.8, 2.8, 6.0],
    "Severity": [2.5, 2.0, 3.0, 8.5],
    "Kriteria quality gate": [6.0, 3.5, 5.0, 2.5],
    "Tindakan": [2.0, 7.0, 3.2, 3.8],
    "TE": [1.2, 9.0, 1.6, 2.0, 2.2],
}
for table in doc.tables:
    if table.rows:
        header = table.rows[0].cells[0].text.strip()
        if header in TABLE_WIDTHS:
            set_widths(table, TABLE_WIDTHS[header])

for p in doc.paragraphs:
    txt = p.text.strip()
    # table and figure captions
    if txt.startswith("Tabel 7."):
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        for r in p.runs:
            r.font.size = Pt(9)
            r.font.bold = True
    elif txt.startswith("Gambar 7."):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        for r in p.runs:
            r.font.size = Pt(9)
            r.font.italic = True
    # page break before main sections
    if txt == "Ringkasan Eksekutif":
        p.paragraph_format.page_break_before = True
    # center and size images
    if p._p.xpath('.//w:drawing'):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        for shape in p._p.xpath('.//wp:inline'):
            # width element in EMU
            for ext in shape.xpath('./wp:extent'):
                cx = int(ext.get('cx')); cy = int(ext.get('cy'))
                target = int(Cm(16).emu)
                if cx > target:
                    width_el = shape.xpath('./wp:extent')[0]
                    width_el.set('cx', str(target))
                    width_el.set('cy', str(int(cy * target / cx)))
                # also update a:ext in graphic
                for aext in shape.xpath('.//a:ext'):
                    aext.set('cx', str(target if cx > target else cx))
                    aext.set('cy', str(int(cy * target / cx) if cx > target else cy))

doc.save(DOCX)
print("docx saved")
